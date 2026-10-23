"""Реплики и тексты игры -> таблица для редактора (английский текст правят в колонке «New English»).
python art_tests/dialogs_export.py  -> art_tests/dialogs/WheelsAndBones_dialogues.xlsx
Обратно: python art_tests/dialogs_import.py <файл.xlsx>  (правки -> i18n_work/en9_fixes15_editor.json).

Каждая строка привязана к квесту, варианту поломки и шагу (колонки Quest / Variant / Where / Step), чтобы редактор
видел, к чему относится реплика: подсказка Барти идёт рядом с номером и текстом шага, у ловушек пометка WRONG hint.
Лист «Job texts» — весь сценарий квеста по порядку: улики, диагнозы, шаги с заметками, мини-игры, факт, проезд.
Одинаковые строки (например «Включи пульт.») выводятся один раз — в игре это одна и та же строка."""
import os, sys
os.environ.setdefault("SDL_VIDEODRIVER", "dummy")
os.environ.setdefault("SDL_AUDIODRIVER", "dummy")
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, ".."))
import pygame
pygame.init()
pygame.display.set_mode((1280, 720))
import game
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

SRC = open(os.path.join(HERE, "..", "game.py"), encoding="utf-8").read()
OUT = os.path.join(HERE, "dialogs", "WheelsAndBones_dialogues.xlsx")
game.set_lang("en")
EN = game.tr

HEAD = ["Quest", "Variant", "Where in the game", "Step", "Who / what", "English now", "New English", "Comment", "key"]
rows = {"Dialogues": [], "Story": [], "Job texts": []}
seen = set()
SPEAKERS = {"Бренна", "Барти", "Бартоломью", "Мартин"}


def add(sheet, quest, variant, where, step, who, ru, repeat=False):
    """repeat=True: повтор строки показать для контекста (без ключа — правится первое вхождение)."""
    if not ru or not isinstance(ru, str):
        return
    if ru in seen:
        if repeat:
            rows[sheet].append((quest, variant, where, step, EN(who) if who else "", EN(ru), None))
        return
    seen.add(ru)
    en = EN(ru)
    if who is None:                                   # говорящий — из перевода: «Zix: “…”» -> Zix
        who = en.split(":", 1)[0].strip() if ":" in en[:40] else ""
        SPEAKERS.add(ru.split(":", 1)[0].strip())
    else:
        who = EN(who) if who else ""
    rows[sheet].append((quest, variant, where, step, who, en, ru))


def part_name(pid):
    return EN(game.PARTS.get(pid, pid))


def variants_of(job):
    """[(подпись варианта, словарь с шагами)] — у заданий с вариантами поломки подпись = верный диагноз."""
    if not job.get("variants"):
        return [("", job)]
    out = []
    for n, var in enumerate(job["variants"]):
        ok = next((d["text"] for d in var.get("diagnosis", []) if d.get("correct")), "")
        out.append((f"{'ABCD'[n]} — {EN(ok)}" if ok else "ABCD"[n], var))
    return out


titles = {j["id"]: j["title"] for j in game.ALL_JOBS}
order = [j["id"] for j in game.ALL_JOBS]


def quest_label(job):
    return f"{EN(job['title'])} ({EN(job['client'])})"


# 1. Реплики клиентов и наставников по заданиям — в порядке игры
KIND = {"intro": "client arrives", "worry": "client chats while you work", "learn": "client shares a tip",
        "mistake": "client after your mistake", "good": "client after a good step", "test": "client before the test ride"}
for jid in order:
    talk = game.CLIENT_TALK.get(jid, {})
    job = next(j for j in game.ALL_JOBS if j["id"] == jid)
    q = quest_label(job)
    if job.get("intro"):
        add("Dialogues", q, "", "job start (Guild)", "", None, job["intro"])
    for kind, lines in talk.items():
        for i, ln in enumerate(lines, 1):
            add("Dialogues", q, "", KIND.get(kind, kind), f"line {i} of {len(lines)}", None, ln)
    for vlabel, var in variants_of(job):
        steps = var.get("steps", [])
        for i, st in enumerate(steps, 1):
            if st.get("ghost"):
                where = ("Barty's WRONG hint at a step (trap — the player must NOT follow it)" if st.get("ghost_wrong")
                         else "Barty's hint at a step")
                add("Dialogues", q, vlabel, where, f"Step {i}/{len(steps)}: {EN(st['instr'])}", None, st["ghost"])
        if var.get("thanks"):
            add("Dialogues", q, vlabel, "thanks at the end", "", None, var["thanks"])
for kind, lines in game.BARTY_TALK.items():
    for i, ln in enumerate(lines, 1):
        add("Dialogues", "— any quest —", "", f"Barty's jokes — {kind}", f"line {i} of {len(lines)}", None, ln)
for kind, lines in game.MARTIN_TALK.items():
    for i, ln in enumerate(lines, 1):
        add("Dialogues", "— any quest —", "", f"Martin after a repair — {kind}", f"line {i} of {len(lines)}", None, ln)
for i, ln in enumerate(game.LEG_GAGS, 1):
    add("Dialogues", "— any quest —", "", "Leg gags after a repair", f"line {i} of {len(game.LEG_GAGS)}", None, ln)
# 2. Сюжет: вступление
g = game.Game(headless=True)
for i, (title, text) in enumerate(g.intro_slides):
    add("Story", "Intro", "", f"Intro slide {i + 1} — title", "", "", title)
    add("Story", "Intro", "", f"Intro slide {i + 1} — text", "", "", text)
# прочие реплики «Имя: «…»» из словаря (финал, Гильдия, итоги)
for ru in game._TR:
    head = ru.split(":", 1)[0].strip() if ":" in ru[:30] else ""
    if head in SPEAKERS and "«" in ru and ru[:40] in SRC:      # только то, что ещё есть в игре
        add("Dialogues", "— other scenes —", "", "Other scenes (ending, Guild, results)", "", None, ru)
# 3. Тексты заданий: весь сценарий квеста по порядку
for jid in order:
    job = next(j for j in game.ALL_JOBS if j["id"] == jid)
    q = quest_label(job)
    add("Job texts", q, "", "title", "", "title", job["title"])
    add("Job texts", q, "", "client", "", "client", job["client"])
    add("Job texts", q, "", "problem (story)", "", "problem (story)", job["symptom"])
    if isinstance(job.get("fact"), str):
        add("Job texts", q, "", "what you learned (shown at the end)", "", "fact", job["fact"])
    for vlabel, var in variants_of(job):
        for pid, clue in (var.get("inspect") or {}).items():
            add("Job texts", q, vlabel, "clue when inspecting a part", part_name(pid), "clue", clue)
        for d in var.get("diagnosis", []):
            ok = d.get("correct")
            add("Job texts", q, vlabel, "diagnosis option", "correct" if ok else "wrong", "diagnosis", d["text"])
            if d.get("feedback"):
                add("Job texts", q, vlabel, "why this diagnosis is wrong", EN(d["text"]), "feedback", d["feedback"])
        steps = var.get("steps", [])
        for i, st in enumerate(steps, 1):
            sl = f"Step {i}/{len(steps)}"
            add("Job texts", q, vlabel, "step instruction", sl, "instruction", st.get("instr"), repeat=True)
            if st.get("note"):
                add("Job texts", q, vlabel, "note shown after the step", sl, "note", st["note"], repeat=True)
            if st.get("wrongmsg"):
                add("Job texts", q, vlabel, "message on a wrong move", sl, "wrong move", st["wrongmsg"])
            if st.get("failnote"):
                add("Job texts", q, vlabel, "safety failure explanation (Guild)", sl, "fail note", st["failnote"])
            mini = st.get("mini") or {}
            spec = st.get("spec") or {}
            for src in (mini, spec):
                for k, lab in (("prompt", "mini-game question"), ("sub", "mini-game note"), ("hint", "mini-game hint"),
                               ("label", "mini-game label"), ("reading", "mini-game reading"),
                               ("fb_lo", "mini-game feedback (too low)"), ("fb_hi", "mini-game feedback (too high)")):
                    if isinstance(src.get(k), str):
                        add("Job texts", q, vlabel, lab, sl, "mini-game", src[k])
                for opt in src.get("options", []) or []:
                    text, ok, fb = (list(opt) + ["", ""])[:3]
                    add("Job texts", q, vlabel, "mini-game answer", sl, "correct" if ok else "wrong", text)
                    if fb:
                        add("Job texts", q, vlabel, "mini-game answer — why wrong", sl, "feedback", fb)
        if isinstance(var.get("test"), str):
            add("Job texts", q, vlabel, "test ride", "", "test ride", var["test"])
        if var is not job and isinstance(var.get("fact"), str):
            add("Job texts", q, vlabel, "what you learned (shown at the end)", "", "fact", var["fact"])

wb = Workbook()
info = wb.active
info.title = "How to edit"
help_lines = [
    "Wheels & Bones — dialogue editing sheet",
    "",
    "1. Edit ONLY the column \"New English\". Leave it empty if the line is fine.",
    "2. Columns Quest / Variant / Where / Step say where the line appears: the quest (its title and client), the "
    "breakdown variant (A, B — some quests have two possible causes), the moment in the game and the step number with "
    "its instruction. Barty's hints sit next to the step they belong to; \"WRONG hint\" is a trap the player must refuse.",
    "3. Keep the speaker at the start in the same form: Name: “text”. The game splits the name off at the first colon.",
    "4. Keep lines short: speech bubbles are small (about 2–3 lines on screen). 150 characters is a good maximum.",
    "5. Facts about wheelchairs (numbers, part names, safety) come from the manufacturers' manuals — if you change "
    "a fact, please add a note in \"Comment\" so we can double-check it.",
    "6. Identical lines (e.g. \"Turn the controller on.\") are listed once — the game uses one line everywhere.",
    "7. Do not change the hidden column \"key\" — it links the row to the game.",
    "8. Send the file back (or share the Google Sheet) — the changes go into the game with one command.",
]
for i, t in enumerate(help_lines, 1):
    info.cell(i, 1, t).font = Font(bold=(i == 1), size=14 if i == 1 else 11)
    info.cell(i, 1).alignment = Alignment(wrap_text=True, vertical="top")
info.column_dimensions["A"].width = 130

head_fill = PatternFill("solid", fgColor="FFE9B0")
edit_fill = PatternFill("solid", fgColor="FFFFF4CC")
quest_fills = [PatternFill("solid", fgColor="FFF3F3F3"), PatternFill("solid", fgColor="FFFFFFFF")]
for name, data in rows.items():
    ws = wb.create_sheet(name)
    ws.append(HEAD)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = head_fill
    last_q, band = None, 0
    for quest, variant, where, step, who, en, ru in data:
        if quest != last_q:
            band ^= 1
            last_q = quest
        ws.append([quest, variant, where, step, who, en, "", "" if ru else "(same line as earlier — edit it there)",
                   ru or ""])
        for c in ws[ws.max_row][:6]:
            c.fill = quest_fills[band]
        if not ru:
            ws[ws.max_row][6].fill = PatternFill("solid", fgColor="FFE6E6E6")
    for col, w in zip("ABCDEFGHI", (30, 22, 30, 36, 14, 62, 62, 28, 40)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if row[8].value:
            row[6].fill = edit_fill
    ws.column_dimensions["I"].hidden = True
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = ws.dimensions
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print(OUT, {k: len(v) for k, v in rows.items()})
