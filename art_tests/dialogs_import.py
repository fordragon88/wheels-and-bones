"""Правки редактора из таблицы (art_tests/dialogs_export.py) -> i18n_work/en9_fixes99z_editor.json.
python art_tests/dialogs_import.py <файл.xlsx | файл.csv из Google Sheets> [--dry]
Берёт строки, где заполнена колонка «New English»; ключ — скрытая колонка «key» (русский исходник).
Файл правок сортируется после остальных en*.json, поэтому перекрывает старый перевод. Потом:
python i18n_work/build_lang.py, selftest и проверки — как после любой правки текста."""
import csv, json, os, sys

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "..", "i18n_work", "en9_fixes99z_editor.json")   # сортируется ПОСЛЕ en9_fixes99_feedback — правки редактора главнее
src = sys.argv[1]
rows = []
if src.lower().endswith(".csv"):
    with open(src, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))
else:
    from openpyxl import load_workbook
    wb = load_workbook(src, read_only=True)
    for ws in wb.worksheets:
        rows += [list(r) for r in ws.iter_rows(values_only=True)]
edits, warn = {}, []
head = None
for r in rows:
    r = [("" if c is None else str(c)) for c in r]
    if "New English" in r and "key" in r:
        head = {name: i for i, name in enumerate(r)}
        continue
    if not head or len(r) <= max(head.values()):
        continue
    new, key, old = r[head["New English"]].strip(), r[head["key"]].strip(), r[head["English now"]].strip()
    if not new or not key or new == old:
        continue
    if ":" in old[:40] and ":" not in new[:40]:
        warn.append(f"нет «Имя:» в начале: {new[:70]}")
    if len(new) > 220:
        warn.append(f"длинно ({len(new)}): {new[:70]}")
    edits[key] = new
print(f"правок: {len(edits)}")
for w in warn:
    print("  ВНИМАНИЕ:", w)
if "--dry" not in sys.argv and edits:
    cur = json.load(open(OUT, encoding="utf-8")) if os.path.exists(OUT) else {}
    cur.update(edits)
    json.dump(cur, open(OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("записано:", OUT)
